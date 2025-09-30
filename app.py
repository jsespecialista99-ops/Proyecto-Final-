import streamlit as st
import pandas as pd
import json
from utils.checklist_loader import load_checklist
from utils.document_parser import extract_text_from_files
from fpdf import FPDF

st.set_page_config(page_title="Evaluación Fase Precontractual", layout="wide")
st.markdown(
    """
    <style>
    body, .stApp {
        background-color: #1565c0 !important;
        color: #fff !important;
    }
    .stMarkdown, .stTitle, .stHeader, .stDataFrame, .stTable, .stSubheader {
        color: #fff !important;
    }
    .css-1d391kg, .css-1v0mbdj, .css-1c7y2kd, .css-1v0mbdj, .css-1c7y2kd {
        color: #fff !important;
    }
    /* Botón Browse files */
    [data-testid="stFileUploaderDropzone"] button {
        color: #111 !important;
    }
    /* Botón Descargar informe en PDF */
    [data-testid="stDownloadButton"] button {
        color: #111 !important;
    }
    /* Barra de herramientas Streamlit */
    [data-testid="stToolbar"] {
        color: #111 !important;
    }
    [data-testid="stToolbar"] * {
        color: #111 !important;
    }
    /* Barra de progreso Streamlit */
    [data-testid="stProgressBar"] div[role="progressbar"] {
        background: #fff !important;
    }
    </style>
    """,
    unsafe_allow_html=True
)

# Título principal centrado y en letra blanca
st.markdown(
    """
    <h1 style='text-align: center; color: #fff; font-size: 2.5em; font-weight: bold; margin-bottom: 0.5em; text-transform: uppercase;'>
        CTI ADMINISTRACIÓN PÚBLICA MEDELLÍN
    </h1>
    """,
    unsafe_allow_html=True
)
st.title("Evaluación Fase Precontractual")


# 1. Subida de múltiples documentos con fondo blanco y letra azul
st.markdown(
    """
    <div style='background-color: #fff; color: #1565c0; padding: 10px; border-radius: 8px; font-size: 20px; font-weight: bold; margin-bottom: 10px; text-align:center;'>
        Sube los documentos del contrato (PDF, Word, Excel, CSV)
    </div>
    """,
    unsafe_allow_html=True
)
uploaded_files = st.file_uploader(
    "Sube los documentos del contrato",
    type=["pdf", "docx", "xlsx", "xls", "csv"],
    accept_multiple_files=True,
    label_visibility="visible"
)
st.markdown(
    """
    <style>
    [data-testid="stFileUploaderDropzone"] {
        background: #fff !important;
        color: #1565c0 !important;
        border-radius: 8px;
    }
    [data-testid="stFileUploaderDropzone"] button {
        background: #fff !important;
        color: #1565c0 !important;
        font-weight: bold;
        border-radius: 8px;
    }
    </style>
    """,
    unsafe_allow_html=True
)

# 2. Cargar checklist
checklist = load_checklist("./data/requisitos.json")

# 3. Procesar documentos y evaluar requisitos
results = []
if uploaded_files:
    progress_text = "Procesando documentos..."
    my_bar = st.progress(0, text=progress_text)
    docs_text = {}
    total_files = len(uploaded_files)
    for idx, file in enumerate(uploaded_files):
        # Procesar cada archivo y actualizar barra
        if file.name.lower().endswith('.pdf'):
            from PyPDF2 import PdfReader
            reader = PdfReader(file)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif file.name.lower().endswith('.docx'):
            from docx import Document
            import io
            doc = Document(io.BytesIO(file.read()))
            text = "\n".join([p.text for p in doc.paragraphs])
        elif file.name.lower().endswith('.xlsx') or file.name.lower().endswith('.xls'):
            import pandas as pd
            import io
            excel_file = pd.read_excel(io.BytesIO(file.read()))
            text = excel_file.to_string(index=False)
        elif file.name.lower().endswith('.csv'):
            import pandas as pd
            import io
            csv_file = pd.read_csv(io.BytesIO(file.read()))
            text = csv_file.to_string(index=False)
        else:
            text = ""
        docs_text[file.name] = text
        percent_complete = int(((idx + 1) / total_files) * 100)
        my_bar.progress(percent_complete / 100, text=f"Procesando documentos... {percent_complete}%")

    for req in checklist:
        estado = "no cumple"
        fuente = ""
        for doc_name, doc_text in docs_text.items():
            if req["item"].lower() in doc_text.lower() or req["descripcion"].lower() in doc_text.lower():
                estado = "cumple"
                fuente = doc_name
                break
        results.append({
            "Requisito": req["item"],
            "Descripción": req["descripcion"],
            "Estado": estado,
            "Fuente": fuente
        })

    # Mostrar nombres de archivos cargados en blanco
    st.markdown(
        """
        <div style='color: #fff; font-size: 16px; margin-bottom: 10px;'>
            <b>Documentos cargados:</b><br>
            {} 
        </div>
        """.format("<br>".join([f.name for f in uploaded_files])), unsafe_allow_html=True
    )
    # 4. Mostrar tabla de resultados
    df = pd.DataFrame(results)
    st.subheader("Informe de Evaluación")
    st.dataframe(df)

    # 5. Resumen general
    estados = [r["Estado"] for r in results]
    if all(e == "cumple" for e in estados):
        resumen = "cumple"
    elif any(e == "cumple" for e in estados):
        resumen = "cumple parcialmente"
    else:
        resumen = "no cumple"
    st.markdown(f"**Resumen general:** El contrato {resumen}")

    # 6. Exportar informe en PDF
    def export_pdf(df, resumen):
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.cell(200, 10, txt="Informe de Evaluación de Contrato", ln=True, align="C")
        pdf.cell(200, 10, txt=f"Resumen general: {resumen}", ln=True, align="L")
        pdf.ln(10)
        for idx, row in df.iterrows():
            pdf.multi_cell(0, 10, txt=f"Requisito: {row['Requisito']}\nDescripción: {row['Descripción']}\nEstado: {row['Estado']}\nFuente: {row['Fuente']}\n", border=0)
            pdf.ln(2)
        return pdf.output(dest="S").encode("latin-1")

    pdf_bytes = export_pdf(df, resumen)
    st.download_button("Descargar informe en PDF", data=pdf_bytes, file_name="informe_evaluacion.pdf", mime="application/pdf")
else:
    st.markdown(
        """
        <div style='background-color: #fff; color: #1565c0; padding: 15px; border-radius: 8px; font-size: 18px; font-weight: bold; margin-top: 20px; text-align: center;'>
            Sube los documentos del contrato para iniciar la evaluación
        </div>
        """,
        unsafe_allow_html=True
    )
